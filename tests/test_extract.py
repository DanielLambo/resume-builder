"""Extraction unit tests (txt/tex/docx/pdf + error paths)."""
import pytest

from app.services import extract as extract_mod
from tests.conftest import SAMPLE_LATEX


@pytest.mark.asyncio
async def test_extract_txt_ok(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("Hello resume", encoding="utf-8")
    result = await extract_mod.extract_text(str(p), "a.txt")
    assert result["success"] is True
    assert result["text"] == "Hello resume"


@pytest.mark.asyncio
async def test_extract_txt_empty(tmp_path):
    p = tmp_path / "empty.txt"
    p.write_text("   \n", encoding="utf-8")
    result = await extract_mod.extract_text(str(p), "empty.txt")
    assert result["success"] is False
    assert "empty" in result["error"].lower()


@pytest.mark.asyncio
async def test_extract_tex_ok(tmp_path):
    p = tmp_path / "r.tex"
    p.write_text(SAMPLE_LATEX, encoding="utf-8")
    result = await extract_mod.extract_text(str(p), "r.tex")
    assert result["success"] is True
    assert "\\documentclass" in result["latex"]


@pytest.mark.asyncio
async def test_extract_tex_missing_documentclass(tmp_path):
    p = tmp_path / "bad.tex"
    p.write_text("just words", encoding="utf-8")
    result = await extract_mod.extract_text(str(p), "bad.tex")
    assert result["success"] is False
    assert "documentclass" in result["error"].lower()


@pytest.mark.asyncio
async def test_extract_unsupported():
    result = await extract_mod.extract_text("/tmp/x", "x.exe")
    assert result["success"] is False
    assert "Unsupported" in result["error"]


@pytest.mark.asyncio
async def test_extract_pdf_missing_pdftotext(monkeypatch, tmp_path):
    p = tmp_path / "x.pdf"
    p.write_bytes(b"%PDF-1.4")
    monkeypatch.setattr(extract_mod.shutil, "which", lambda _n: None)
    result = await extract_mod.extract_text(str(p), "x.pdf")
    assert result["success"] is False
    assert "pdftotext" in result["error"].lower()


@pytest.mark.asyncio
async def test_extract_pdf_ok(monkeypatch, tmp_path):
    p = tmp_path / "x.pdf"
    p.write_bytes(b"%PDF-1.4")
    monkeypatch.setattr(extract_mod.shutil, "which", lambda _n: "/usr/bin/pdftotext")

    class FakeProc:
        async def communicate(self):
            return (b"Jane Doe\nEngineer", b"")

    async def fake_exec(*_a, **_k):
        return FakeProc()

    monkeypatch.setattr(extract_mod.asyncio, "create_subprocess_exec", fake_exec)
    result = await extract_mod.extract_text(str(p), "x.pdf")
    assert result["success"] is True
    assert "Jane Doe" in result["text"]


@pytest.mark.asyncio
async def test_extract_pdf_scanned_empty(monkeypatch, tmp_path):
    p = tmp_path / "scan.pdf"
    p.write_bytes(b"%PDF")
    monkeypatch.setattr(extract_mod.shutil, "which", lambda _n: "/usr/bin/pdftotext")

    class FakeProc:
        async def communicate(self):
            return (b"   ", b"")

    async def fake_exec(*_a, **_k):
        return FakeProc()

    monkeypatch.setattr(extract_mod.asyncio, "create_subprocess_exec", fake_exec)
    result = await extract_mod.extract_text(str(p), "scan.pdf")
    assert result["success"] is False
    assert "image" in result["error"].lower() or "scanned" in result["error"].lower()


@pytest.mark.asyncio
async def test_extract_docx_ok(tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    p = tmp_path / "r.docx"
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Software Engineer")
    doc.save(str(p))
    result = await extract_mod.extract_text(str(p), "r.docx")
    assert result["success"] is True
    assert "Jane Doe" in result["text"]


@pytest.mark.asyncio
async def test_extract_docx_empty(tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    p = tmp_path / "empty.docx"
    Document().save(str(p))
    result = await extract_mod.extract_text(str(p), "empty.docx")
    assert result["success"] is False
