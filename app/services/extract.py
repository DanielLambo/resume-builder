import asyncio
import tempfile
from pathlib import Path


async def extract_text(file_path: str, filename: str) -> dict:
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return await _extract_pdf(file_path)
    elif ext == ".docx":
        return _extract_docx(file_path)
    elif ext == ".txt":
        return _extract_txt(file_path)
    elif ext == ".tex":
        return _extract_tex(file_path)
    else:
        return {"success": False, "error": f"Unsupported file type: {ext}. Use PDF, DOCX, TXT, or TEX."}


async def _extract_pdf(file_path: str) -> dict:
    pdftotext = "/opt/homebrew/bin/pdftotext"
    if not Path(pdftotext).exists():
        import shutil
        pdftotext = shutil.which("pdftotext") or ""

    if not pdftotext or not Path(pdftotext).exists():
        return {"success": False, "error": "pdftotext not found. Install with: brew install poppler"}

    try:
        proc = await asyncio.create_subprocess_exec(
            pdftotext, "-layout", file_path, "-",
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        stdout, stderr = await proc.communicate()
        text = stdout.decode("utf-8", errors="replace").strip()
        if not text:
            return {"success": False, "error": "Could not extract text from PDF. It may be image-based/scanned."}
        return {"success": True, "text": text}
    except Exception as e:
        return {"success": False, "error": f"PDF extraction failed: {str(e)}"}


def _extract_docx(file_path: str) -> dict:
    try:
        from docx import Document
        doc = Document(file_path)
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        text = "\n".join(paragraphs)
        if not text:
            return {"success": False, "error": "Could not extract text from DOCX. File may be empty."}
        return {"success": True, "text": text}
    except Exception as e:
        return {"success": False, "error": f"DOCX extraction failed: {str(e)}"}


def _extract_txt(file_path: str) -> dict:
    try:
        text = Path(file_path).read_text(encoding="utf-8", errors="replace").strip()
        if not text:
            return {"success": False, "error": "Text file is empty."}
        return {"success": True, "text": text}
    except Exception as e:
        return {"success": False, "error": f"Text extraction failed: {str(e)}"}


def _extract_tex(file_path: str) -> dict:
    try:
        text = Path(file_path).read_text(encoding="utf-8", errors="replace").strip()
        if not text:
            return {"success": False, "error": "TeX file is empty."}
        if "\\documentclass" not in text:
            return {"success": False, "error": "File doesn't look like a valid .tex file (missing \\documentclass)."}
        return {"success": True, "latex": text}
    except Exception as e:
        return {"success": False, "error": f"TeX extraction failed: {str(e)}"}
